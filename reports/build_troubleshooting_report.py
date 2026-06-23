from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out_dir = Path('reports')
out_dir.mkdir(exist_ok=True)
out_path = out_dir / 'TaxiFlow_Troubleshooting_Report_2026-05-19.docx'

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.12
for style_name, size, color in [
    ('Title', 22, RGBColor(23, 43, 77)),
    ('Heading 1', 15, RGBColor(23, 43, 77)),
    ('Heading 2', 12, RGBColor(31, 88, 122)),
    ('Heading 3', 10.5, RGBColor(23, 43, 77)),
]:
    st = styles[style_name]
    st.font.name = 'Arial'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True
    st.paragraph_format.space_before = Pt(10 if style_name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(5)
styles['Subtitle'].font.name = 'Arial'
styles['Subtitle']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
styles['Subtitle'].font.size = Pt(11)
styles['Subtitle'].font.color.rgb = RGBColor(91, 103, 112)

def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(rows, widths=None):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, text in enumerate(rows[0]):
        set_cell_text(hdr[i], text, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(hdr[i], '172B4D')
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

def number(text):
    doc.add_paragraph(text, style='List Number')

def callout(title, body, fill='EEF6F8'):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(23, 43, 77)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('TaxiFlow Pro Troubleshooting Report')
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(23, 43, 77)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('Authentication, Manager/Admin Access, and Production Deployment Recovery')
r2.font.size = Pt(11)
r2.font.color.rgb = RGBColor(91, 103, 112)
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.add_run('Prepared: 19 May 2026 | Workspace: C:\\Users\\appri\\taxiflowv2').italic = True

doc.add_paragraph()
callout('Current operating position', 'The live app is usable after switching production from the broken Supabase authentication endpoint to local setup authentication and after correcting Admin/Manager module defaults. Admin and Manager users can now access operational modules. The remaining known issue is cosmetic/data-retention related: the seeded Manager display name may revert to Lerato Maseko after edits.', 'E8F3EC')

doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph('TaxiFlow Pro experienced two separate but related access problems. First, the live sign-in page failed with a generic "Failed to fetch" message while the app was configured for Supabase authentication. Second, after local authentication was restored, Admin and Manager users could sign in but were limited to viewing basic screens because their module access defaults only enabled Overview and Settings.')
doc.add_paragraph('The Supabase issue was traced to an unreachable production Supabase project hostname. The module access issue was traced to a hard-coded default in the application runtime. The live environment was adjusted to use local setup authentication, and the code was patched so Admin and Manager users receive operational module access by default.')
add_table([
    ['Area', 'Status', 'Notes'],
    ['Live login', 'Resolved for local setup auth', 'Production now shows Local setup authentication instead of Supabase authentication.'],
    ['Admin access', 'Resolved', 'Admin receives operational modules allowed by role.'],
    ['Manager access', 'Resolved for module access', 'Manager receives operational modules allowed by role.'],
    ['Manager name persistence', 'Known remaining issue', 'Seeded fallback name Lerato Maseko can reappear. Not blocking operations.'],
    ['Supabase-backed auth', 'Deferred', 'Current Supabase hostname does not resolve. Needs correct project URL/key before re-enabling.'],
], widths=[1.7, 1.7, 4.7])

doc.add_heading('2. Incident Timeline', level=1)
add_table([
    ['Sequence', 'Finding / Action', 'Outcome'],
    ['1', 'Live app displayed Supabase authentication and returned Failed to fetch on owner login.', 'Confirmed this was not a wrong-password error. It was a network/auth-provider fetch failure.'],
    ['2', 'Production bundle was inspected for Supabase configuration.', 'Production embedded VITE_SUPABASE_URL=https://vfxftviwteypasredjgm.supabase.co and a publishable key.'],
    ['3', 'Direct DNS/connectivity check was performed against the Supabase host.', 'Host could not be resolved, explaining the browser Failed to fetch behavior.'],
    ['4', 'Vercel production environment variables were changed to placeholder Supabase values.', 'App no longer initializes Supabase auth and falls back to local setup authentication.'],
    ['5', 'Production was redeployed.', 'New deployment succeeded and login screen switched to Local setup authentication.'],
    ['6', 'Admin/Manager access was tested after login.', 'They could sign in, but modules remained locked/view-only due to application defaults.'],
    ['7', 'Runtime role defaults were patched and pushed to main.', 'Admin and Manager now receive Money, Fleet & Operations, Drivers, and Settings where allowed by role.'],
], widths=[0.8, 3.6, 3.7])

doc.add_heading('3. Root Cause Analysis', level=1)
doc.add_heading('3.1 Supabase authentication failure', level=2)
doc.add_paragraph('The production app was configured to initialize Supabase authentication, but the configured Supabase project hostname did not resolve. In the browser this surfaced as a generic "Failed to fetch" message on sign-in.')
bullet('Production URL observed: https://vfxftviwteypasredjgm.supabase.co')
bullet('Behavior: authentication failed before password validation could complete.')
bullet('Impact: all users relying on Supabase authentication were locked out of the live app.')

doc.add_heading('3.2 Admin and Manager module access failure', level=2)
doc.add_paragraph('Once local authentication was restored, Admin and Manager users could authenticate. However, the runtime default module access function only enabled Overview and Settings for those roles. Money, Fleet & Operations, and Drivers were part of the role catalog but were not enabled by default.')
bullet('File changed: src/lib/appRuntime.js')
bullet('Function changed: createDefaultModuleViewAccess(role)')
bullet('Previous behavior: Admin and Manager defaulted to overview + settings only.')
bullet('Corrected behavior: Admin and Manager default to every module listed as allowed for their role.')

doc.add_heading('3.3 Manager name reverting to Lerato Maseko', level=2)
doc.add_paragraph('The Manager display name remains a known non-blocking issue. The app has a seeded account directory that includes manager@taxiflow.local as Lerato Maseko. If the saved appUsers record is refreshed, reset, or not persisted as expected, the seeded fallback can reappear.')
bullet('Likely source: AUTH_ACCOUNT_DIRECTORY and default app user rebuilding in src/lib/appRuntime.js.')
bullet('Impact: cosmetic/user-profile persistence issue only; it does not currently block module access or daily operations.')
bullet('Recommended future fix: separate seeded defaults from editable user profiles and make saved appUsers the single source of truth for display names.')

doc.add_heading('4. Fixes Implemented', level=1)
doc.add_heading('4.1 Production authentication recovery', level=2)
doc.add_paragraph('Vercel production environment variables were changed so the app does not attempt to initialize the unreachable Supabase project. This restores local setup authentication for live operations.')
add_table([
    ['Variable', 'Current production value / behavior'],
    ['VITE_BACKEND_MODE', 'live'],
    ['VITE_SUPABASE_URL', 'Placeholder URL, causing Supabase auth to be disabled by app config guard.'],
    ['VITE_SUPABASE_ANON_KEY', 'Placeholder key, causing Supabase auth to be disabled by app config guard.'],
], widths=[2.4, 5.7])

doc.add_heading('4.2 Role access correction', level=2)
doc.add_paragraph('The runtime access default now uses the module role catalog as the source of truth. This means if MODULE_VIEW_ACCESS says Admin or Manager can access a module, the default module access state now enables that module for that role.')
add_table([
    ['Commit', 'Purpose'],
    ['d3adf04', 'Fix admin and manager module access'],
    ['1851932', 'Clarify manager access settings copy'],
], widths=[1.5, 6.6])

doc.add_heading('5. Verification Completed', level=1)
bullet('Production build completed successfully after the role access patch.')
bullet('Changes were committed and pushed to origin/main.')
bullet('Vercel reported the latest production deployment as Ready.')
bullet('The deployed bundle was checked and confirmed to contain the access-fix wording.')
bullet('User observation after deployment: Admin and Manager users now have access to all operational features.')

doc.add_heading('6. Current Operating Instructions', level=1)
number('Open the app using the production URL and include ?reset-sw=1 if any browser still shows stale behavior.')
number('Use local setup authentication while Supabase remains disabled.')
number('Owner login: owner@taxiflow.local with password TaxiFlow.123.')
number('Admin login: admin@taxiflow.local with password TaxiFlow.123.')
number('Manager login: manager@taxiflow.local with password TaxiFlow.123.')
number('If users still see stale screens, sign out, open the reset URL, and sign in again.')
callout('Cache reset URL', 'https://taxiflowpro.vercel.app/?reset-sw=1', 'FFF5D6')

doc.add_heading('7. Deployment Procedure Going Forward', level=1)
doc.add_paragraph('For code changes, run build before staging and committing. Avoid git add . when there may be generated files, test artifacts, local env files, or drive-specific files.')
add_table([
    ['Purpose', 'Recommended command sequence'],
    ['Code update', 'git status; npm run build; git add <specific files>; git commit -m "message"; git push origin HEAD:main'],
    ['Vercel env/settings only', 'Update Vercel env; redeploy a specific deployment or run a new production deploy; verify live URL.'],
    ['Service worker/cache issue', 'Open https://taxiflowpro.vercel.app/?reset-sw=1'],
], widths=[2.1, 6.0])

doc.add_heading('8. Risks and Recommended Next Steps', level=1)
doc.add_heading('Immediate risks', level=2)
bullet('Local setup authentication is operational but not a permanent enterprise-grade authentication model.')
bullet('Any user with the known local password can sign in as seeded accounts if they know the email/password combination.')
bullet('Manager display name persistence remains unresolved and may confuse users, even though functionality works.')
doc.add_heading('Recommended next steps', level=2)
number('Decide whether Supabase Auth should be restored. If yes, confirm the correct Supabase project URL and anon key from the Supabase dashboard.')
number('Add SUPABASE_SERVICE_ROLE_KEY to Vercel only if server-side user creation through /api/admin/auth-users is required. Never expose it as a VITE_ variable.')
number('Refactor user creation so live Supabase Auth user creation happens before appUsers snapshot creation. This prevents local-only users that cannot sign in.')
number('Fix manager name persistence by treating saved appUsers as authoritative and only applying seeded defaults when no saved profile exists.')
number('Add regression tests for Owner, Admin, and Manager module access so this does not regress.')

doc.add_heading('Appendix A: Key URLs and References', level=1)
add_table([
    ['Item', 'Value'],
    ['Production app', 'https://taxiflowpro.vercel.app/'],
    ['Cache reset URL', 'https://taxiflowpro.vercel.app/?reset-sw=1'],
    ['Latest production deployment observed', 'https://taxiflowpro-gr0tgaqpl-apprigate.vercel.app'],
    ['Repository remote', 'https://github.com/Tebza102/taxiflowpro.git'],
    ['Workspace path', 'C:\\Users\\appri\\taxiflowv2'],
], widths=[2.3, 5.8])

doc.add_heading('Appendix B: Files Changed for Access Fix', level=1)
bullet('src/lib/appRuntime.js - changed Admin/Manager default module access logic.')
bullet('src/App.jsx - updated Settings panel text to reflect operational access defaults.')

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.text = 'TaxiFlow Pro Troubleshooting Report | 19 May 2026'
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(91, 103, 112)

doc.save(out_path)
print(out_path.resolve())
